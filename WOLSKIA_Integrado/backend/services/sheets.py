"""Stateless spreadsheet operations. No uploads are persisted or sent to an LLM."""
import base64
import csv
import io
import json
import math
import re
import zipfile
import time
from collections import deque
from threading import Lock, BoundedSemaphore
from copy import copy
from datetime import date, datetime
from flask import Blueprint, request, jsonify
from openpyxl import Workbook, load_workbook

sheets = Blueprint('sheets', __name__)
MAX_BYTES = 8 * 1024 * 1024
MAX_CELLS = 1_000_000
MAX_COLUMNS = 1024
SIZE_LIMIT_MESSAGE = "La planilla supera el límite de 1.000.000 de celdas o 1.024 columnas."
_rate = {}
_lock = Lock()
_slots = BoundedSemaphore(2)


def load(data, name):
    if not data or len(data) > MAX_BYTES:
        raise ValueError('El archivo debe contener datos y pesar como máximo 8 MB.')
    if name.lower().endswith('.csv'):
        text = data.decode('utf-8-sig')
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=',;\t')
        except csv.Error:
            dialect = csv.excel
        book = Workbook()
        for row in csv.reader(io.StringIO(text), dialect):
            if book.active.max_row * max(1, len(row)) > MAX_CELLS or len(row) > MAX_COLUMNS:
                raise ValueError(SIZE_LIMIT_MESSAGE)
            book.active.append(row)
            for cell in book.active[book.active.max_row]:
                # CSV text is never promoted to executable Excel formulas.
                if cell.data_type == 'f':
                    cell.data_type = 's'
    elif name.lower().endswith('.xlsx'):
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 40 * 1024 * 1024:
                raise ValueError('El contenido descomprimido supera 40 MB.')
        book = load_workbook(io.BytesIO(data), data_only=False, keep_links=False)
    else:
        raise ValueError('Esta vista admite XLSX y CSV UTF-8. Para XLS usá las herramientas existentes.')
    if sum(s.max_row * s.max_column for s in book) > MAX_CELLS or any(s.max_column > MAX_COLUMNS for s in book):
        raise ValueError(SIZE_LIMIT_MESSAGE)
    return book


def display(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        return None
    return value if isinstance(value, (str, int, float, bool, type(None))) else str(value)


def column(headers, name):
    if headers.count(name) != 1:
        raise ValueError('Indicá un encabezado exacto y único. Disponibles: ' + ', '.join(headers))
    return headers.index(name) + 1


def parse_number(value):
    if isinstance(value, str) and value.startswith('='):
        raise ValueError('No se calculan fórmulas. Seleccioná una columna de valores.')
    try:
        result = float(value)
    except (ValueError, TypeError):
        raise ValueError('Hay valores no numéricos o vacíos; revisá la columna.')
    if not math.isfinite(result):
        raise ValueError('Hay valores numéricos inválidos.')
    return result


def as_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    for pattern in ('%Y-%m-%d', '%d/%m/%Y'):
        try:
            return datetime.strptime(str(value).strip(), pattern).date()
        except ValueError:
            pass
    raise ValueError('Las fechas deben usar AAAA-MM-DD o DD/MM/AAAA, o ser fechas reales de Excel.')


def derived_sheet(book, ws, rows, title):
    if ws.merged_cells.ranges or any(c.data_type == 'f' for row in ws for c in row):
        raise ValueError('Para reordenar o filtrar, usá una hoja sin fórmulas ni celdas combinadas; así se evitan referencias incorrectas.')
    target = book.create_sheet(title)
    for output_index, source_index in enumerate([1, *rows], 1):
        for source in ws[source_index]:
            cell = target.cell(output_index, source.column, source.value)
            cell._style = copy(source._style)
            if source.data_type == 's':
                cell.data_type = 's'
    for key, dim in ws.column_dimensions.items():
        target.column_dimensions[key].width = dim.width
    return target


def operate(book, sheet, command):
    ws = book[sheet]
    headers = [str(c.value or '') for c in ws[1]]
    command = command.strip()
    if not command:
        return 'Vista previa del archivo original.', False
    match = re.fullmatch(r'(?:sumar|sumá|suma) (.+) por (.+)', command, re.I)
    if match:
        value_name, group_name = match.groups()
        vi, gi = column(headers, value_name), column(headers, group_name)
        groups = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            key = str(display(row[gi - 1]) or '(sin grupo)')
            groups.setdefault(key, []).append(parse_number(row[vi - 1]))
        if len(groups) > 100:
            raise ValueError('El resumen supera 100 grupos. Reducí el archivo.')
        return '\n'.join(f'{key}: {math.fsum(values):g}' for key, values in groups.items()), False
    match = re.fullmatch(r'(?:filtrar|filtrá) (.+) entre (\S+) y (\S+)', command, re.I)
    if match:
        name, start, end = match.groups()
        ci = column(headers, name)
        start, end = as_date(start), as_date(end)
        if start > end:
            raise ValueError('La fecha inicial debe ser anterior a la final.')
        rows = [r for r in range(2, ws.max_row + 1) if start <= as_date(ws.cell(r, ci).value) <= end]
        derived_sheet(book, ws, rows, 'SURI filtro')
        return f'{len(rows)} registros entre {start} y {end}. Resultado en una hoja nueva; el original se conserva.', True
    match = re.fullmatch(r'(?:ordenar|ordená) por (.+)', command, re.I)
    if match:
        ci = column(headers, match[1])
        values = [ws.cell(r, ci).value for r in range(2, ws.max_row + 1)]
        if any(v is None for v in values):
            raise ValueError('La columna contiene celdas vacías. Completalas antes de ordenar.')
        try:
            keys = [as_date(v) for v in values]
        except ValueError:
            if all(isinstance(v, (float, int)) for v in values):
                keys = values
            elif all(isinstance(v, str) for v in values):
                keys = [v.casefold() for v in values]
            else:
                raise ValueError('La columna mezcla tipos de datos. Revisá sus valores.')
        rows = [r for _, r in sorted(zip(keys, range(2, ws.max_row + 1)))]
        derived_sheet(book, ws, rows, 'SURI ordenado')
        return f'{len(rows)} registros ordenados en una hoja nueva. El original se conserva.', True
    if command.lower() in ('duplicados', 'buscar duplicados', 'buscá duplicados'):
        seen, duplicates = set(), []
        for index, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
            key = tuple(row)
            if key in seen:
                duplicates.append(index)
            seen.add(key)
        return f'Filas duplicadas: {duplicates[:100]}. Total: {len(duplicates)}.', False
    match = re.fullmatch(r'(?:sumar|sumá|suma) (.+)', command, re.I)
    if match:
        name = match[1].strip()
        if headers.count(name) != 1:
            raise ValueError('Indicá un encabezado exacto y único. Disponibles: ' + ', '.join(headers))
        values = [r[0] for r in ws.iter_rows(min_row=2, min_col=headers.index(name)+1, max_col=headers.index(name)+1, values_only=True)]
        if any(isinstance(v, str) and v.startswith('=') for v in values):
            raise ValueError('La columna contiene fórmulas. No se dispone de un motor de recálculo.')
        numbers = []
        for value in values:
            if value is None or value == '':
                continue
            try:
                number = float(value)
            except (ValueError, TypeError):
                raise ValueError('La columna contiene valores no numéricos; no se estimará el resultado.')
            if not math.isfinite(number):
                raise ValueError('La columna contiene valores numéricos inválidos.')
            numbers.append(number)
        return f'Suma de {name}: {math.fsum(numbers):g} · {len(numbers)} valores.', False
    match = re.fullmatch(r'(?:cambiar|cambiá) ([A-Z]+[1-9][0-9]*) (?:por|a) (.+)', command, re.I)
    if match:
        coordinate, raw = match.groups()
        cell = ws[coordinate.upper()]
        if cell.row > ws.max_row or cell.column > ws.max_column:
            raise ValueError('La celda debe estar dentro de la tabla existente.')
        if cell.data_type == 'f':
            raise ValueError('No se sobrescriben fórmulas desde esta herramienta.')
        value = json.loads(raw) if raw[:1] in ('"', '-', '0','1','2','3','4','5','6','7','8','9') else raw
        if not isinstance(value, (str, int, float)) or isinstance(value, bool):
            raise ValueError('Usá un texto o un número.')
        if isinstance(value, float) and not math.isfinite(value):
            raise ValueError('El número no es válido.')
        if isinstance(value, str) and value.startswith(('=', '+', '-', '@')):
            raise ValueError('No se aceptan fórmulas en texto nuevo.')
        before = display(cell.value)
        cell.value = value
        return f'{coordinate.upper()}: {before!r} → {value!r}. Se descargará una copia.', True
    raise ValueError('Usá «sumar Ingresos», «sumar Ingresos por Persona», «buscar duplicados», «cambiar B2 por 125», «ordenar por Fecha» o «filtrar Fecha entre 2026-10-01 y 2026-10-08». Los encabezados deben coincidir exactamente.')


@sheets.post('/api/ai/sheets')
def process():
    # Bound multipart parsing even when Content-Length is absent.
    request.max_content_length = MAX_BYTES + 64_000
    if request.content_length and request.content_length > MAX_BYTES + 64_000:
        return jsonify(message='El archivo supera 8 MB.'), 413
    upload = request.files.get('file')
    if upload is None:
        return jsonify(message='Falta adjuntar un archivo.'), 400
    identity = request.headers.get('X-Wolskia-Client-IP') or request.remote_addr or 'local'
    with _lock:
        now = time.monotonic()
        for key in list(_rate):
            if not _rate[key] or now - _rate[key][-1] >= 60:
                del _rate[key]
        bucket = _rate.setdefault(identity[:64], deque())
        while bucket and now - bucket[0] >= 60:
            bucket.popleft()
        if len(bucket) >= 20 or len(_rate) > 1024:
            return jsonify(message='Demasiadas operaciones. Esperá un minuto.'), 429
        bucket.append(now)
    if not _slots.acquire(blocking=False):
        return jsonify(message='Hay otras planillas en proceso. Reintentá en unos segundos.'), 429
    try:
        book = load(upload.stream.read(MAX_BYTES + 1), upload.filename or '')
        sheet = request.form.get('sheet') or book.sheetnames[0]
        if sheet not in book.sheetnames:
            raise ValueError('La hoja seleccionada no existe.')
        source_sheet, source_names = sheet, list(book.sheetnames)
        original_sheet_count = len(book.sheetnames)
        summary, changed = operate(book, sheet, request.form.get('command', ''))
        if len(book.sheetnames) > original_sheet_count:
            sheet = book.sheetnames[-1]
        ws = book[sheet]
        result = dict(sheets=source_names, sheet=source_sheet, preview_sheet=sheet, rows=ws.max_row, columns=ws.max_column,
                      preview=[[display(v) for v in row] for row in ws.iter_rows(max_row=min(ws.max_row, 51), values_only=True)],
                      summary=summary, changed=changed,
                      warning='XLSX: se conservan estilos y fórmulas compatibles; no se recalculan. Macros, vínculos externos y objetos avanzados no están garantizados. CSV se convierte a XLSX.')
        if changed:
            buffer = io.BytesIO()
            book.save(buffer)
            output = buffer.getvalue()
            if len(output) > MAX_BYTES:
                raise ValueError('El archivo generado supera 8 MB.')
            check = load_workbook(io.BytesIO(output))
            if check.sheetnames != book.sheetnames:
                raise ValueError('Falló la validación del archivo generado.')
            result['download'] = base64.b64encode(output).decode('ascii')
        response = jsonify(result)
        response.headers['Cache-Control'] = 'no-store'
        return response
    except (ValueError, KeyError, UnicodeError, zipfile.BadZipFile, csv.Error) as error:
        return jsonify(message=str(error)), 400
    except Exception:
        return jsonify(message='No se pudo interpretar la planilla. Verificá su formato y que no esté protegida.'), 400
    finally:
        _slots.release()
