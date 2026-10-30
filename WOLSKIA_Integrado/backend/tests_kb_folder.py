"""Pruebas offline con documentos sintéticos, sin datos de la empresa."""
import importlib.util
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from kb.store import Store
from kb.folder import ingest_folder
from kb.documents import fragments, pdf_chunks, DependencyMissing

READERS = bool(importlib.util.find_spec('pypdf') and importlib.util.find_spec('docx'))

def write_pdf(path):
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject,NameObject,DecodedStreamObject
    w=PdfWriter(); p=w.add_blank_page(width=300,height=300)
    font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
    p[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):font})})
    stream=DecodedStreamObject(); stream.set_data(b'BT /F1 12 Tf 30 200 Td (Recepcion sintetica del taller) Tj ET')
    p[NameObject('/Contents')]=w._add_object(stream)
    w.add_blank_page(width=300,height=300)
    with path.open('wb') as f:w.write(f)

def write_docx(path, text='Verificar orden ficticia'):
    from docx import Document
    w=Document();w.add_heading('Entrega',level=1);w.add_paragraph(text)
    w.add_table(rows=1,cols=1).cell(0,0).text='Tabla sintetica'
    w.save(path)

class BasicFolderTests(unittest.TestCase):
    def test_fragmentos_con_solapamiento_y_pagina(self):
        rows=list(fragments(' '.join('palabra'+str(i) for i in range(1700)),'s.pdf','pdf','Titulo',4))
        self.assertEqual(rows[0]['page'],4)
        self.assertEqual(rows[0]['text'].split()[-100:],rows[1]['text'].split()[:100])
        self.assertIn('palabra1699',rows[-1]['text'])
    def test_dependencias_ausentes_no_rompen_servidor(self):
        with patch('kb.documents.optional',return_value=None):
            with self.assertRaises(DependencyMissing):list(pdf_chunks(Path('no.pdf'),'no.pdf',[]))

@unittest.skipUnless(READERS,'Lectores opcionales: corré en el entorno offline.')
class FolderTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory(); self.root=Path(self.tmp.name)/'manual';self.root.mkdir()
        self.store=Store(Path(self.tmp.name)/'kb.sqlite');self.logs=[]
    def tearDown(self):
        self.store.close();self.tmp.cleanup()
    def run_folder(self,**kw):return ingest_folder(self.root,self.store,log=self.logs.append,**kw)
    def test_pdf_pagina_ocr_y_docx_titulo_tabla(self):
        write_pdf(self.root/'ejemplo.pdf');write_docx(self.root/'ejemplo.docx')
        result=self.run_folder();self.assertEqual(result['indexados'],2)
        rows=self.store.search('recepcion');self.assertEqual(rows[0]['page'],1)
        self.assertEqual(rows[0]['path_relativo'],'manual/ejemplo.pdf')
        self.assertTrue(self.store.search('Tabla sintetica'))
        doc=self.store.db.execute("SELECT etapa FROM chunks WHERE source_type='docx'").fetchone()
        self.assertEqual(doc[0],'Entrega')
        issues=self.store.db.execute('SELECT page,kind FROM ingestion_issues').fetchall()
        self.assertEqual([(r[0],r[1]) for r in issues],[(2,'requiere_ocr')])
    def test_duplicados_por_contenido_y_reanudacion(self):
        write_docx(self.root/'a.docx'); (self.root/'sub').mkdir()
        (self.root/'sub'/'copia.docx').write_bytes((self.root/'a.docx').read_bytes())
        first=self.run_folder();self.assertEqual(first['duplicados'],1)
        count=self.store.db.execute('SELECT count(*) FROM chunks').fetchone()[0]
        with patch('kb.folder.docx_chunks',side_effect=AssertionError('No debe reprocesar')):
            again=self.run_folder()
        self.assertEqual(again['reanudados'],1)
        self.assertEqual(self.store.db.execute('SELECT count(*) FROM chunks').fetchone()[0],count)
    def test_pdf_duplicado_tambien_con_paginas_pendientes_ocr(self):
        write_pdf(self.root/'a.pdf')
        (self.root/'b.pdf').write_bytes((self.root/'a.pdf').read_bytes())
        result=self.run_folder();self.assertEqual(result['duplicados'],1)
        self.assertEqual(self.store.db.execute("SELECT count(*) FROM chunks WHERE source_type='pdf'").fetchone()[0],1)

    def test_actualizacion_reemplaza_y_no_duplica(self):
        path=self.root/'a.docx';write_docx(path,'Textoanterior');self.run_folder()
        write_docx(path,'Textonuevo');self.run_folder()
        self.assertEqual(self.store.search('Textoanterior'),[])
        self.assertTrue(self.store.search('Textonuevo'))
    def test_error_un_archivo_no_frena_otros_y_registra(self):
        (self.root/'a.pdf').write_bytes(b'PDF roto')
        write_docx(self.root/'b.docx')
        result=self.run_folder();self.assertEqual(result['omitidos'],1);self.assertEqual(result['indexados'],1)
        self.assertTrue(self.store.db.execute("SELECT 1 FROM ingestion_issues WHERE kind='error_archivo'").fetchone())
    def test_interrupcion_conserva_terminados_y_revierte_actual(self):
        write_docx(self.root/'a.docx');write_docx(self.root/'b.docx','Otro contenido')
        from kb.documents import docx_chunks
        def interrupt(path,source,issues):
            if path.name=='b.docx':
                yield dict(source_file=source,text='Incompleto')
                raise KeyboardInterrupt()
            yield from docx_chunks(path,source,issues)
        with patch('kb.folder.docx_chunks',side_effect=interrupt):
            with self.assertRaises(KeyboardInterrupt):self.run_folder()
        self.assertTrue(self.store.db.execute("SELECT 1 FROM ingestion_files WHERE source='manual/a.docx'").fetchone())
        self.assertEqual(self.store.search('Incompleto'),[])
        result=self.run_folder();self.assertEqual(result['reanudados'],1);self.assertEqual(result['indexados'],1)
    def test_ocr_no_se_ejecuta_por_defecto(self):
        write_pdf(self.root/'a.pdf')
        import kb.documents as documents
        real=documents.optional
        def guard(name):
            if name in ('pytesseract','PIL.Image'):raise AssertionError('OCR no autorizado')
            return real(name)
        with patch('kb.documents.optional',side_effect=guard):self.run_folder()
        self.assertTrue(any('requiere_ocr' in line for line in self.logs))
    def test_ocr_flag_cambia_perfil_y_reintenta(self):
        write_pdf(self.root/'a.pdf');self.run_folder()
        result=self.run_folder(ocr=True)
        self.assertEqual(result['indexados'],1)
        self.assertTrue(any('OCR hacen falta' in line for line in self.logs))

if __name__ == '__main__':unittest.main()
