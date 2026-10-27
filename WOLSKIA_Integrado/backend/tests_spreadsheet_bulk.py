import base64
import io
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from services.spreadsheet_chat import apply_plan, bounds, prefix_plan, plan, snapshot


class BulkSpreadsheetTests(unittest.TestCase):
    def book(self, count=55):
        book = Workbook()
        book.remove(book.active)
        for i in range(count):
            ws = book.create_sheet(f'Hoja {i + 1}')
            ws.append(['Teléfono', 'Nombre'])
            ws.append([5491512345678, 'Ejemplo'])
            ws.append(['+5491598765432', 'Otro'])
            ws.append(['11549150000', 'Sin cambios'])
            ws.append(['=5491512345678', 'Fórmula'])
            ws['A2'].font = Font(bold=True)
        return book

    def test_one_action_replaces_all_55_sheets_and_preserves_formulas(self):
        result = apply_plan(self.book(), {'actions': [{'op': 'replace_prefix', 'sheet': '*', 'from': '54915', 'to': '54911'}]})
        output = load_workbook(io.BytesIO(base64.b64decode(result['download'])))
        self.assertEqual(len(output.sheetnames), 55)
        self.assertIn('110 celdas', result['summary'])
        for ws in output:
            self.assertEqual(ws['A2'].value, 5491112345678)
            self.assertEqual(ws['A3'].value, '+5491198765432')
            self.assertEqual(ws['A4'].value, '11549150000')
            self.assertEqual(ws['A5'].value, '=5491512345678')
            self.assertTrue(ws['A2'].font.bold)

    def test_grouped_header_action_does_not_need_55_model_actions(self):
        output = apply_plan(self.book(), {'actions': [{'op': 'style_header', 'sheet': '*', 'color': '112233'}]})
        self.assertEqual(output['status'], 'ok')

    def test_more_than_50_distinct_actions_allowed(self):
        book = self.book()
        result = apply_plan(book, {'actions': [{'op': 'style_header', 'sheet': ws.title, 'color': '112233'} for ws in book]})
        self.assertEqual(result['status'], 'ok')

    def test_limits_still_apply(self):
        book = self.book(257)
        with self.assertRaises(ValueError):
            bounds(book)

    def test_scope_list_and_column(self):
        book = self.book(2)
        result = apply_plan(book, {'actions': [{'op': 'replace_prefix', 'sheet': ['Hoja 1'], 'column': 'Teléfono', 'from': '54915', 'to': '54911'}]})
        output = load_workbook(io.BytesIO(base64.b64decode(result['download'])))
        self.assertEqual(output['Hoja 1']['A2'].value, 5491112345678)
        self.assertEqual(output['Hoja 2']['A2'].value, 5491512345678)

    def test_no_matches_does_not_report_success(self):
        with self.assertRaises(ValueError):
            apply_plan(self.book(1), {'actions': [{'op': 'replace_prefix', 'sheet': '*', 'from': '888', 'to': '999'}]})

    def test_unsafe_prefix_rejected(self):
        with self.assertRaises(ValueError):
            apply_plan(self.book(1), {'actions': [{'op': 'replace_prefix', 'sheet': '*', 'from': '54915', 'to': '=1'}]})

    def test_model_not_called_for_explicit_all_numbers(self):
        with patch('services.ai_service.check_local_rate_limit', return_value=(True, '')), patch('services.ai_service.get_providers') as providers:
            spec = plan('que cambie todos los numeros de 54915 a 54911', [], self.book(), 'test')
            self.assertEqual(spec['actions'][0]['sheet'], '*')
            providers.assert_not_called()

    def test_ambiguous_or_scoped_request_not_guessed(self):
        for request in ['Cambiá todos los 54915 por 54911 excepto Pepe', 'Cambiá solo la hoja 1 de 54915 a 54911', 'Cambiá todos los números de 54915 a 54911 y sumá 3', 'No cambies todos los numeros de 54915 a 54911', 'Explicame como cambiar todos los numeros de 54915 a 54911']:
            self.assertIsNone(prefix_plan(request, self.book(1)))

    def test_preview_gives_later_sheets_a_share(self):
        book = self.book()
        for i in range(200):
            book.worksheets[0].append(['dato extenso ' * 10, 'otra fila' * 10])
        result = snapshot(book)
        self.assertTrue(result[-1]['preview'])
        self.assertLess(len(str(result)), 28000)


if __name__ == '__main__':
    unittest.main()
