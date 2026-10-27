import base64
import io
import unittest
from pathlib import Path
from unittest.mock import patch

from flask import Flask
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.table import Table

from services.spreadsheet_chat import apply_plan, load_input, process_chat, snapshot, value


CREATE = {'filename': 'trabajadores.xlsx', 'sheets': [{'name': 'Trabajadores', 'rows': [
    ['Trabajador'], ['Papa'], ['Pepe'], ['Pipi'], ['Popo']]}]}


class SpreadsheetChatTests(unittest.TestCase):
    def setUp(self):
        app = Flask(__name__)
        app.add_url_rule('/api/ai/sheets/chat', view_func=lambda: process_chat('test'), methods=['POST'])
        self.client = app.test_client()
        self.book = Workbook()
        self.book.active.title = 'Personal'
        self.book.active.append(['Trabajador', 'Horas', 'Total'])
        self.book.active.append(['Papa', 8, '=B2*2'])
        self.book.active.append(['Pepe', 6, None])
        self.book.active['B2'].font = Font(bold=True, color='FF0000')
        self.book.create_sheet('Otra')['A1'] = 'Sin cambios'

    def bytes(self):
        stream = io.BytesIO()
        self.book.save(stream)
        return stream.getvalue()

    def unpack(self, result):
        return load_workbook(io.BytesIO(base64.b64decode(result['download'])))

    def edit(self, actions):
        return apply_plan(self.book, {'filename': 'copia.xlsx', 'actions': actions})

    def test_create_requested_workers(self):
        result = apply_plan(None, CREATE)
        book = self.unpack(result)
        self.assertEqual(list(book.active.values), [('Trabajador',), ('Papa',), ('Pepe',), ('Pipi',), ('Popo',)])
        self.assertEqual(book.active.freeze_panes, 'A2')
        self.assertTrue(book.active['A1'].font.bold)
        self.assertEqual(result['filename'], 'trabajadores.xlsx')

    def test_creation_with_numeric_values_and_formula(self):
        result = apply_plan(None, {'sheets': [{'name': 'Horas', 'rows': [['Nombre', 'Horas'], ['Papa', 8], ['Pepe', 6], ['Total', '=SUM(B2:B3)']]}]})
        self.assertEqual(self.unpack(result).active['B4'].value, '=SUM(B2:B3)')

    def test_edit_preserves_other_sheet_styles_formulas_and_original(self):
        original = self.bytes()
        result = self.edit([{'op': 'update_rows', 'sheet': 'Personal', 'where': {'column': 'Trabajador', 'equals': 'papa'}, 'values': {'Horas': 10}}])
        book = self.unpack(result)
        self.assertEqual(book['Personal']['B2'].value, 10)
        self.assertTrue(book['Personal']['B2'].font.bold)
        self.assertEqual(book['Personal']['C2'].value, '=B2*2')
        self.assertEqual(book['Otra']['A1'].value, 'Sin cambios')
        self.assertEqual(load_workbook(io.BytesIO(original))['Personal']['B2'].value, 8)

    def test_followup_uses_latest_output(self):
        result = self.edit([{'op': 'add_column', 'sheet': 'Personal', 'header': 'Turno', 'default': 'Mañana'}])
        book = self.unpack(result)
        result = apply_plan(book, {'actions': [{'op': 'append_rows', 'sheet': 'Personal', 'rows': [['Pipi', 7, None, 'Tarde']]}]})
        book = self.unpack(result)
        self.assertEqual(book['Personal']['D1'].value, 'Turno')
        self.assertEqual(book['Personal']['A4'].value, 'Pipi')
        self.assertEqual(book['Personal']['D4'].value, 'Tarde')

    def test_native_excel_table_extends_with_new_column_and_row(self):
        self.book['Personal'].add_table(Table(displayName='PersonalTable', ref='A1:C3'))
        self.book = load_workbook(io.BytesIO(self.bytes()))
        result = self.edit([
            {'op': 'add_column', 'sheet': 'Personal', 'header': 'Turno', 'default': 'Mañana'},
            {'op': 'append_rows', 'sheet': 'Personal', 'rows': [['Pipi', 7, None, 'Tarde']]},
        ])
        book = self.unpack(result)
        table = book['Personal'].tables['PersonalTable']
        self.assertEqual(table.ref, 'A1:D4')
        self.assertEqual([c.name for c in table.tableColumns], ['Trabajador', 'Horas', 'Total', 'Turno'])

    def test_delete_and_sort_simple_rows(self):
        self.book['Personal']['C2'] = None
        result = self.edit([
            {'op': 'delete_rows', 'sheet': 'Personal', 'where': {'column': 'Trabajador', 'equals': 'Papa'}},
            {'op': 'append_rows', 'sheet': 'Personal', 'rows': [['Ana', 7, None]]},
            {'op': 'sort_rows', 'sheet': 'Personal', 'column': 'Trabajador'},
        ])
        self.assertEqual(self.unpack(result)['Personal']['A2'].value, 'Ana')

    def test_unknown_and_duplicate_columns_rejected(self):
        for action in [
            {'op': 'add_column', 'sheet': 'Personal', 'header': 'Horas'},
            {'op': 'update_rows', 'sheet': 'Personal', 'where': {'column': 'Inexistente', 'equals': 'Papa'}, 'values': {'Horas': 5}},
        ]:
            with self.assertRaises(ValueError):
                self.edit([action])

    def test_set_cell_can_clear_value(self):
        result = self.edit([{'op': 'set_cells', 'sheet': 'Personal', 'cells': {'B2': None, 'B3': 9}}])
        self.assertIsNone(self.unpack(result)['Personal']['B2'].value)
        self.assertEqual(self.unpack(result)['Personal']['B3'].value, 9)

    def test_update_row_can_clear_value(self):
        result = self.edit([{'op': 'update_rows', 'sheet': 'Personal', 'where': {'column': 'Trabajador', 'equals': 'Pepe'}, 'values': {'Horas': None}}])
        self.assertIsNone(self.unpack(result)['Personal']['B3'].value)

    def test_legacy_xls_edit_and_conversion(self):
        data = (Path(__file__).parent / 'tests/fixtures/personal-legacy.xls').read_bytes()
        book, warning = load_input(data, 'personal.xls')
        result = apply_plan(book, {'actions': [{'op': 'update_rows', 'sheet': 'Personal', 'where': {'column': 'Trabajador', 'equals': 'Pepe'}, 'values': {'Horas': 9}}]})
        self.assertEqual(self.unpack(result)['Personal']['B3'].value, 9)
        self.assertTrue(result['filename'].endswith('.xlsx'))
        self.assertIn('últimos valores', warning)

    def test_refuse_fake_xls(self):
        with self.assertRaises(ValueError):
            load_input(self.bytes(), 'mentira.xls')

    def test_missing_person_does_not_return_fake_success(self):
        with self.assertRaises(ValueError):
            self.edit([{'op': 'update_rows', 'sheet': 'Personal', 'where': {'column': 'Trabajador', 'equals': 'Nadie'}, 'values': {'Horas': 10}}])

    def test_unsupported_code_not_executed(self):
        with self.assertRaises(ValueError):
            self.edit([{'op': 'python', 'sheet': 'Personal', 'code': 'import os'}])

    def test_reorder_with_formulas_rejected(self):
        with self.assertRaises(ValueError):
            self.edit([{'op': 'sort_rows', 'sheet': 'Personal', 'column': 'Trabajador'}])

    def test_rename_with_formula_rejected(self):
        with self.assertRaises(ValueError):
            self.edit([{'op': 'rename_sheet', 'sheet': 'Personal', 'name': 'Otro'}])

    def test_formula_network_and_external_refs_rejected(self):
        for formula in ['=HYPERLINK("https://example.com")', '=WEBSERVICE(A1)', "='[otra.xlsx]Hoja1'!A1", '=INDIRECT(A1)', '=cmd|calc!A1', '=SUM(A1:A3)+EVIL()']:
            with self.subTest(formula=formula), self.assertRaises(ValueError):
                value(formula)

    def test_sheet_and_dimensions_limited(self):
        for spec in [{'sheets': [{'name': 'No/valido', 'rows': [['A']]}]}, {'sheets': [{'name': 'Hoja', 'rows': [['A'] * 1025]}]}]:
            with self.assertRaises(ValueError):
                apply_plan(None, spec)

    def test_snapshot_is_bounded_and_marks_truncation(self):
        for i in range(70):
            self.book.active.append([str(i), 1, None])
        view = snapshot(self.book)
        self.assertTrue(view[0]['truncated'])
        self.assertLessEqual(len(view[0]['preview']), 51)

    @patch('services.spreadsheet_chat.plan', return_value=CREATE)
    def test_endpoint_returns_valid_download(self, planner):
        response = self.client.post('/api/ai/sheets/chat', data={'message': 'Hacé un Excel con Papa, Pepe, Pipi y Popo'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers['Cache-Control'], 'no-store')
        self.assertEqual(self.unpack(response.json).active['A5'].value, 'Popo')

    @patch('services.spreadsheet_chat.plan', return_value={'question': '¿Qué columnas querés agregar?'})
    def test_clarification_has_no_fake_download(self, planner):
        response = self.client.post('/api/ai/sheets/chat', data={'message': 'Creá un Excel'})
        self.assertEqual(response.json['status'], 'clarification')
        self.assertNotIn('download', response.json)

    @patch('services.spreadsheet_chat.plan', return_value={'actions': [{'op': 'set_cells', 'sheet': 'Personal', 'cells': {'B2': 12}}]})
    def test_multipart_edit_download(self, planner):
        original = (Path(__file__).parent / 'tests/fixtures/personal-original.xlsx').read_bytes()
        response = self.client.post('/api/ai/sheets/chat', data={'message': 'Cambiá a 12 las horas de Papa', 'file': (io.BytesIO(original), 'original.xlsx')})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(self.unpack(response.json)['Personal']['B2'].value, 12)

    def test_bad_input_fails_before_provider(self):
        with patch('services.spreadsheet_chat.plan') as planner:
            for data in [{'message': ''}, {'message': 'x' * 4001}, {'message': 'Editá', 'history': '{}'}, {'message': 'Editá', 'file': (io.BytesIO(b'not excel'), 'a.xls')}]:
                self.assertEqual(self.client.post('/api/ai/sheets/chat', data=data).status_code, 400)
            planner.assert_not_called()


if __name__ == '__main__':
    unittest.main()
