import base64
import io
import unittest
import zipfile
from flask import Flask
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from services.sheets import sheets, load, operate, _rate


class SheetTests(unittest.TestCase):
    def setUp(self):
        _rate.clear()
        self.book = Workbook()
        self.book.active.append(['Persona', 'Ingresos', 'Total'])
        self.book.active.append(['Ana', 10, '=B2*2'])
        self.book.active.append(['Luis', 25, None])
        self.book.active['B2'].font = Font(bold=True)
        app = Flask(__name__)
        app.register_blueprint(sheets)
        self.client = app.test_client()

    def data(self):
        stream = io.BytesIO()
        self.book.save(stream)
        return stream.getvalue()

    def post(self, **fields):
        return self.client.post('/api/ai/sheets', data={'file': (io.BytesIO(self.data()), 'prueba.xlsx'), **fields})

    def test_preview(self):
        result = self.post().get_json()
        self.assertEqual(result['preview'][1][0], 'Ana')
        self.assertNotIn('download', result)

    def test_sum(self):
        self.assertIn('35', operate(self.book, 'Sheet', 'sumar Ingresos')[0])

    def test_grouped_sum(self):
        self.book.active.append(['Ana', 5, None])
        summary, changed = operate(self.book, 'Sheet', 'sumar Ingresos por Persona')
        self.assertIn('Ana: 15', summary)
        self.assertFalse(changed)

    def test_filter_includes_boundaries(self):
        book = load(b'Fecha,Ingreso\n2026-10-01,10\n2026-10-08,25\n2026-10-09,99', 'a.csv')
        summary, changed = operate(book, 'Sheet', 'filtrar Fecha entre 2026-10-01 y 2026-10-08')
        self.assertTrue(changed)
        self.assertEqual(book.worksheets[-1].max_row, 3)
        self.assertIn('2 registros', summary)

    def test_sort_dates_and_preserve_source(self):
        book = load(b'Fecha,Ingreso\n08/10/2026,25\n01/10/2026,10', 'a.csv')
        operate(book, 'Sheet', 'ordenar por Fecha')
        self.assertEqual(book.worksheets[-1]['A2'].value, '01/10/2026')
        self.assertEqual(book.worksheets[0]['A2'].value, '08/10/2026')

    def test_sort_formula_refused(self):
        with self.assertRaises(ValueError):
            operate(self.book, 'Sheet', 'ordenar por Ingresos')

    def test_ambiguous_header(self):
        self.book.active['C1'] = 'Ingresos'
        with self.assertRaises(ValueError):
            operate(self.book, 'Sheet', 'sumar Ingresos')

    def test_formula_not_estimated(self):
        with self.assertRaises(ValueError):
            operate(self.book, 'Sheet', 'sumar Total')

    def test_duplicates(self):
        self.book.active.append(['Luis', 25, None])
        self.assertIn('Total: 1', operate(self.book, 'Sheet', 'buscar duplicados')[0])

    def test_edit_preserves_original_styles_formulas(self):
        original = self.data()
        response = self.post(command='cambiar B2 por 125')
        self.assertEqual(response.status_code, 200)
        result = load_workbook(io.BytesIO(base64.b64decode(response.json['download'])))
        self.assertEqual(result.active['B2'].value, 125)
        self.assertTrue(result.active['B2'].font.bold)
        self.assertEqual(result.active['C2'].value, '=B2*2')
        self.assertEqual(load_workbook(io.BytesIO(original)).active['B2'].value, 10)

    def test_formula_overwrite_rejected(self):
        self.assertEqual(self.post(command='cambiar C2 por 5').status_code, 400)

    def test_formula_injection_rejected(self):
        self.assertEqual(self.post(command='cambiar B2 por =HYPERLINK("https://example.com")').status_code, 400)

    def test_invalid_and_unsupported(self):
        for content, name in [(b'bad', 'a.xlsx'), (b'bad', 'a.xls'), (b'', 'a.csv')]:
            with self.assertRaises((ValueError, zipfile.BadZipFile)):
                load(content, name)

    def test_csv_formulas_remain_text(self):
        book = load(b'Persona;Ingresos\nAna;10\nLuis;=1+1', 'datos.csv')
        self.assertEqual(book.active['B3'].data_type, 's')

    def test_rate_limit(self):
        for _ in range(20):
            self.post()
        self.assertEqual(self.post().status_code, 429)

    def test_unknown_sheet_and_code(self):
        self.assertEqual(self.post(sheet='No existe').status_code, 400)
        self.assertEqual(self.post(command='import os').status_code, 400)


if __name__ == '__main__':
    unittest.main()
