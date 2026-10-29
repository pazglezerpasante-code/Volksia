import base64
import io
import unittest
import zipfile
from openpyxl import Workbook, load_workbook
from services.xlsx_prefix import process_prefix_xlsx
from services.spreadsheet_chat import bounds


class NativePrefixTests(unittest.TestCase):
    def source(self):
        book = Workbook()
        book.active['A1'] = 'Teléfono'
        book.active['A2'] = '5491512345678'
        book.active['A3'] = 5491598765432
        book.active['A4'] = '=5491512345678'
        book.active['B1'] = 'Sin cambios'
        stream = io.BytesIO()
        book.save(stream)
        return stream.getvalue()

    def output(self, raw):
        result = process_prefix_xlsx(raw, '54915', '54911')
        return result, base64.b64decode(result['download'])

    def test_roundtrip_cells_and_unrelated_package_members(self):
        original = self.source()
        result, output = self.output(original)
        book = load_workbook(io.BytesIO(output))
        self.assertEqual(book.active['A2'].value, '5491112345678')
        self.assertEqual(book.active['A3'].value, 5491198765432)
        self.assertEqual(book.active['A4'].value, '=5491512345678')
        self.assertIn('2 números', result['summary'])
        with zipfile.ZipFile(io.BytesIO(original)) as a, zipfile.ZipFile(io.BytesIO(output)) as b:
            self.assertEqual(a.namelist(), b.namelist())
            for name in a.namelist():
                if name != 'xl/worksheets/sheet1.xml':
                    self.assertEqual(a.read(name), b.read(name))

    def test_shared_strings_are_extended_and_original_entry_retained(self):
        raw = self.source()
        out = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(raw)) as original, zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as target:
            for name in original.namelist():
                data = original.read(name)
                if name == 'xl/worksheets/sheet1.xml':
                    data = data.replace(b'<c r="A2" t="inlineStr"><is><t>5491512345678</t></is></c>', b'<c r="A2" t="s"><v>0</v></c>')
                if name == 'xl/_rels/workbook.xml.rels':
                    data = data.replace(b'</Relationships>', b'<Relationship Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings" Target="sharedStrings.xml" Id="rIdStrings"/></Relationships>')
                target.writestr(name, data)
            target.writestr('xl/sharedStrings.xml', b'<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" count="1" uniqueCount="1"><si><r><rPr><b/></rPr><t>54915</t></r><r><t>12345678</t></r></si></sst>')
        result, output = self.output(out.getvalue())
        with zipfile.ZipFile(io.BytesIO(output)) as verified:
            strings = verified.read('xl/sharedStrings.xml')
            self.assertIn(b'<t>54915</t>', strings)
            self.assertIn(b'<t>54911</t>', strings)
            self.assertIn(b'uniqueCount="2"', strings)
            self.assertIn(b'<c r="A2" t="s"><v>1</v></c>', verified.read('xl/worksheets/sheet1.xml'))

    def test_old_cell_count_no_longer_rejects_250k_area(self):
        book = Workbook()
        book.active['A1'] = 'Dato'
        book.active['C100001'] = 'Dato'
        bounds(book)

    def test_one_million_limit_still_enforced(self):
        book = Workbook()
        book.active['Z100001'] = 'Dato'
        with self.assertRaises(ValueError):
            bounds(book)


if __name__ == '__main__':
    unittest.main()
