"""Lectura offline del instructivo. Las macros nunca se ejecutan ni conservan."""
import io
import re
import posixpath
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET
from openpyxl import load_workbook
from .store import CONTACT, redact

PROCESOS = {'AGENDAMIENTO','RECEPCION','EJECUCION','ENTREGA'}
NS = {'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main',
      'a':'http://schemas.openxmlformats.org/drawingml/2006/main',
      'xdr':'http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing'}

def package_extras(archive, sheet_path):
    """Recupera notas y formas por relaciones de la hoja, sin evaluar contenido."""
    relpath = posixpath.join(posixpath.dirname(sheet_path),'_rels',posixpath.basename(sheet_path)+'.rels')
    notes, shapes = {}, []
    if relpath not in archive.namelist(): return notes, shapes
    for rel in ET.fromstring(archive.read(relpath)):
        if rel.get('TargetMode') == 'External': continue
        target = rel.get('Target','')
        target = target.lstrip('/') if target.startswith('/') else posixpath.normpath(posixpath.join(posixpath.dirname(sheet_path),target))
        if target not in archive.namelist(): continue
        kind = rel.get('Type','').rsplit('/',1)[-1]
        if kind == 'comments':
            for node in ET.fromstring(archive.read(target)).findall('.//s:comment',NS):
                notes[node.get('ref')] = ''.join(t.text or '' for t in node.findall('.//s:t',NS))
        elif kind == 'drawing':
            for anchor in ET.fromstring(archive.read(target)):
                texts = [t.text or '' for t in anchor.findall('.//a:t',NS)]
                if texts:
                    shapes.append((' / '.join(texts),anchor.findtext('xdr:from/xdr:row',default='?',namespaces=NS),anchor.findtext('xdr:from/xdr:col',default='?',namespaces=NS)))
    return notes, shapes

def extract(path):
    raw = Path(path).read_bytes()
    if len(raw) > 8*1024*1024: raise ValueError('El instructivo supera 8 MB.')
    return extract_bytes(raw,Path(path).name)

def extract_bytes(raw, source):
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        if sum(i.file_size for i in archive.infolist()) > 40*1024*1024 or len(archive.infolist()) > 10000:
            raise ValueError('El contenido interno del Excel excede el límite seguro.')
        book = load_workbook(io.BytesIO(raw),keep_vba=False,keep_links=False,data_only=False)
        relations = {r.get('Id'):r.get('Target') for r in ET.fromstring(archive.read('xl/_rels/workbook.xml.rels'))}
        sheet_nodes = ET.fromstring(archive.read('xl/workbook.xml')).findall('s:sheets/s:sheet',NS)
        result = []
        def emit(ws,text,paso=None,responsable='',sin_asociar=False):
            # Cada parte conserva la referencia al mismo paso; no se descartan notas largas.
            sensitive = bool(CONTACT.search(text))
            words = redact(text).split()
            for start in range(0,len(words),650):
                result.append(dict(source_file=source,source_type='xlsm',etapa=ws.title,hoja=ws.title,paso=paso,
                    responsable=redact(responsable),text=' '.join(words[start:start+750]),path_relativo=source,
                    sensible=int(sensitive),sin_asociar=int(sin_asociar)))
        for ws,node in zip(book,sheet_nodes):
            if ws.max_row*ws.max_column > 1000000: raise ValueError('Hoja demasiado grande.')
            rid = node.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
            target = relations[rid]
            sheetpath = target.lstrip('/') if target.startswith('/') else posixpath.normpath('xl/'+target)
            notes,shapes = package_extras(archive,sheetpath)
            cells = [c for row in ws for c in row if c.value is not None]
            for row in ws:
                for c in row:
                    if c.comment: notes.setdefault(c.coordinate,c.comment.text)
            if ws.title.upper() in PROCESOS:
                for coord,note in notes.items():
                    cell = ws[coord]; title = str(cell.value or '')
                    prefix = re.match(r'\s*(\d{1,3})\s*[).:-]',note)
                    # El diagrama real coloca el número a la derecha del título.
                    neighbors = [c for c in cells if c.row==cell.row and 0<c.column-cell.column<=8 and re.fullmatch(r'\d{1,3}',str(c.value).strip())]
                    nearest = min(neighbors,key=lambda c:c.column) if neighbors else None
                    step = int(prefix[1]) if prefix else (int(nearest.value) if nearest is not None and len(title)>4 and title not in ('<--','-->') else None)
                    if prefix and nearest is not None and int(nearest.value)!=step: step=None
                    nearby = []
                    for label in cells:
                        if label.column == cell.column and 0 <= label.row-cell.row <= 3 and re.match(r'^Resp\s*\.?\s*:',str(label.value),re.I):
                            neighbors_resp = [x for x in cells if x.row==label.row and 0 < x.column-label.column <= 5]
                            value = min(neighbors_resp,key=lambda x:x.column).value if neighbors_resp else ''
                            nearby.append(str(label.value)+' '+str(value))
                    resp = re.search(r'Resp\s*\.?\s*:\s*([^\n]+)',note+'\n'+'\n'.join(nearby),re.I)
                    text = f'Hoja: {ws.title}. Paso: {step if step is not None else "sin_asociar"}. Celda: {coord}. Título: {title}\n{note}'
                    if nearby: text += '\n'+'\n'.join(nearby)
                    emit(ws,text,step,resp[1].strip() if resp else '',step is None)
                # Decisiones y textos sin nota se conservan sin inferir flechas ni rutas.
                for c in cells:
                    if c.coordinate not in notes and isinstance(c.value,str) and len(c.value.strip())>1:
                        nums = [x for x in cells if x.row==c.row and 0<x.column-c.column<=8 and re.fullmatch(r'\d{1,3}',str(x.value).strip())]
                        labels = [x for x in cells if x.column==c.column and 0<x.row-c.row<=3 and re.match(r'^Resp\s*\.?\s*:',str(x.value),re.I)]
                        step = int(min(nums,key=lambda x:x.column).value) if nums and labels and len(c.value)>4 else None
                        resp = ''
                        if labels:
                            label = min(labels,key=lambda x:x.row)
                            values = [x for x in cells if x.row==label.row and 0<x.column-label.column<=5]
                            if values: resp = str(min(values,key=lambda x:x.column).value)
                        emit(ws,f'Celda {c.coordinate}: {c.value}. Responsable: {resp}',step,resp,sin_asociar=step is None)
                for text,row,col in shapes:
                    emit(ws,f'Forma en fila {row}, columna {col}: {text}',sin_asociar=True)
            elif ws.title.upper() == 'MENU DE INICIO':
                for c in cells:
                    if re.search(r'autor|fecha|versi[oó]n',str(c.value),re.I):
                        near = [f'{x.coordinate}: {x.value}' for x in cells if x.row==c.row and 0<=x.column-c.column<=3]
                        emit(ws,'Metadatos: '+' / '.join(near))
            else:
                rows = [[c for c in row if c.value is not None] for row in ws]
                candidates = [row for row in rows[:10] if row]
                header = max(candidates,key=lambda row:sum(isinstance(c.value,str) for c in row),default=[])
                heading = ' | '.join(f'{c.column_letter}: {c.value}' for c in header)
                for row in ws:
                    filled = [c for c in row if c.value is not None]
                    if filled: emit(ws,'Encabezado: '+heading+'; Tabla '+ws.title+'; fila '+str(row[0].row)+': '+' | '.join(f'{c.column_letter}: {c.value}' for c in filled))
                for coord,note in notes.items(): emit(ws,f'Nota {coord}: {note}',sin_asociar=True)
        book.close()
        return sorted(result,key=lambda c:(c['etapa'],c['paso'] is None,c['paso'] or 0))
